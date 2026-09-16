

import sys
from pyomo.environ import (
    ConcreteModel, Set, Var, Constraint, Objective,
    NonNegativeReals, NonNegativeIntegers, Binary,
    minimize, SolverFactory, value
)

import data as dt  # data.py must sit in the same folder as this script

try:
    from openpyxl import Workbook
except ImportError:
    print("The 'openpyxl' package is required to write the Excel output files.")
    print("Install it first by running the following command in the VS Code terminal:")
    print("    pip install openpyxl")
    sys.exit(1)

# ----------------------------------------------------------------------------
# Path to the GLPK solver executable on this machine
# ----------------------------------------------------------------------------
solver = SolverFactory('glpk')
# solver.executable = "C:/GLPK/winglpk-4.65/winglpk-4.65/glpk-4.65/w64/glpsol.exe"

# ----------------------------------------------------------------------------
# 1. Basic index sets (derived from the data actually available in data.py)
# ----------------------------------------------------------------------------
PRODUCTS = list(range(1, 29))    # i = 1..28  (q = 28 final products)
MATERIALS = list(range(1, 18))   # j = 1..17  (m = 17 raw materials)
WEEKS = list(range(1, 13))       # t = 1..12  (n = 12 planning weeks)

PROTEIN_MATERIALS = [1, 2, 3]                 # red meat, chicken, turkey
COLD_MATERIALS = [4, 5, 6, 7, 8]              # cheese and vegetables needing refrigeration
DRY_MATERIALS = list(range(9, 18))            # dry / non-perishable materials

# ----------------------------------------------------------------------------
# 2. Data validation and (limited, transparent) automatic repair
#
#    The sales-price table r_it in data.py is incomplete:
#      - product 7  is missing weeks 10, 11, 12
#      - product 18 is missing ALL 12 weeks
#    Everything else in data.py (costs, demand, capacities, formulation
#    coefficients, etc.) was checked and is complete for i=1..28, j=1..17,
#    t=1..12.
#
#    For product 7 there is enough of its own price history (9 out of 12
#    weeks) to extrapolate the missing weeks using its own average weekly
#    price increase. This is done automatically below and clearly reported.
#
#    For product 18 there is NO price data at all, so there is no honest way
#    to guess a selling price for it. The script stops with a clear message
#    instead of inventing a number for a real product.
# ----------------------------------------------------------------------------


def repair_r_it(r_it, products, weeks):
    fixed = dict(r_it)
    auto_filled = {}
    unfixable = {}

    for i in products:
        present = sorted((t, r_it[(i, t)]) for t in weeks if (i, t) in r_it)
        missing = [t for t in weeks if (i, t) not in r_it]
        if not missing:
            continue

        if len(present) >= 2:
            increments = [present[k + 1][1] - present[k][1] for k in range(len(present) - 1)]
            avg_increment = sum(increments) / len(increments)
            last_t, last_val = present[-1]
            filled_here = {}
            still_missing = []
            for t in sorted(missing):
                if t > last_t:
                    last_val = last_val + avg_increment
                    last_t = t
                    fixed[(i, t)] = round(last_val, 2)
                    filled_here[t] = fixed[(i, t)]
                else:
                    still_missing.append(t)
            if filled_here:
                auto_filled[i] = filled_here
            if still_missing:
                unfixable[i] = still_missing
        else:
            unfixable[i] = missing

    return fixed, auto_filled, unfixable


r_it, auto_filled, unfixable = repair_r_it(dt.r_it, PRODUCTS, WEEKS)

if auto_filled:
    print("=" * 78)
    print("NOTICE: r_it (selling price) had missing weeks that were auto-filled")
    print("        by extrapolating that product's own price trend:")
    for i, filled in auto_filled.items():
        weeks_txt = ", ".join(f"t={t} -> {val}" for t, val in filled.items())
        print(f"  product {i}: {weeks_txt}")
    print("=" * 78)

if unfixable:
    print("=" * 78)
    print("ERROR: r_it (selling price) is missing data that cannot be safely")
    print("       auto-filled because the product has little or no price")
    print("       history in data.py:")
    for i, missing in unfixable.items():
        print(f"  product {i}: missing weeks {missing}")
    print("Please add the missing r_it values for these products in data.py")
    print("and run the script again.")
    print("=" * 78)
    sys.exit(1)

# Initial-condition assumption not given anywhere in data.py or the document:
# no product has a delivery delay (backorder) before week 1.
PF_I0 = {i: 0.0 for i in PRODUCTS}

# ----------------------------------------------------------------------------
# 3. Small helper functions for "previous period" references (t-1, t-2, t-3)
# ----------------------------------------------------------------------------


def inv_prev(model, j, t):
    return model.Inv[j, t - 1] if t > 1 else dt.I_j0[j]


def sf_prev(model, i, t):
    return model.sf[i, t - 1] if t > 1 else dt.sf_i0[i]


def pf_prev(model, i, t):
    return model.pf[i, t - 1] if t > 1 else PF_I0[i]


def w_prev(model, t):
    return model.W[t - 1] if t > 1 else dt.W_0


# ----------------------------------------------------------------------------
# 4. Build the Pyomo model
# ----------------------------------------------------------------------------
model = ConcreteModel(name="Andre_GoalProgramming_ProductionPlanning")

model.I = Set(initialize=PRODUCTS, ordered=True)
model.J = Set(initialize=MATERIALS, ordered=True)
model.T = Set(initialize=WEEKS, ordered=True)

# ---------------------------- Decision variables -----------------------------
model.fi = Var(model.I, model.T, domain=NonNegativeReals)   # sales quantity
model.Pn = Var(model.I, model.T, domain=NonNegativeReals)   # normal-time production
model.Po = Var(model.I, model.T, domain=NonNegativeReals)   # overtime production
model.W = Var(model.T, domain=NonNegativeIntegers)          # workforce size
model.H = Var(model.T, domain=NonNegativeIntegers)          # hired workers
model.L = Var(model.T, domain=NonNegativeIntegers)          # laid-off workers
model.X = Var(model.J, model.T, domain=NonNegativeReals)    # raw material purchased
model.Y = Var(model.J, model.T, domain=NonNegativeReals)    # raw material consumed
model.pf = Var(model.I, model.T, domain=NonNegativeReals)   # delivery delay (backorder)
model.Inv = Var(model.J, model.T, domain=NonNegativeReals)  # raw material end-of-period stock
model.sf = Var(model.I, model.T, domain=NonNegativeReals)   # finished product end-of-period stock
model.y = Var(model.I, model.T, domain=Binary)              # 1 if product i is made in period t

# ---------------------------- Auxiliary deviation variables ------------------
model.d_p_plus = Var(domain=NonNegativeReals)
model.d_p_minus = Var(domain=NonNegativeReals)
model.d_cw_plus = Var(domain=NonNegativeReals)
model.d_cw_minus = Var(domain=NonNegativeReals)
model.d_id_plus = Var(domain=NonNegativeReals)
model.d_id_minus = Var(domain=NonNegativeReals)

# ----------------------------------------------------------------------------
# 5. Goal constraints (link Z1, Z2, Z3 to their deviation variables)
# ----------------------------------------------------------------------------


def profit_goal_rule(m):
    revenue = sum(m.fi[i, t] * r_it[i, t] for i in m.I for t in m.T)
    material_cost = sum(dt.c_jt[j, t] * m.X[j, t] for j in m.J for t in m.T)
    normal_labor_cost = sum(dt.cn_t[t] * m.W[t] for t in m.T)
    overtime_labor_cost = sum(
        m.Po[i, t] * dt.co_t[t] * dt.wn_it[i] for i in m.I for t in m.T
    )
    material_holding_cost = sum(dt.f_jt[j, t] * m.Inv[j, t] for j in m.J for t in m.T)
    product_holding_cost = sum(dt.h_it[i, t] * m.sf[i, t] for i in m.I for t in m.T)
    transport_cost = sum(dt.q_it[i, t] * m.fi[i, t] for i in m.I for t in m.T)
    hire_fire_cost = sum(dt.ch_t[t] * m.H[t] + dt.cl_t[t] * m.L[t] for t in m.T)

    profit = (
        revenue
        - material_cost
        - normal_labor_cost
        - overtime_labor_cost
        - material_holding_cost
        - product_holding_cost
        - transport_cost
        - hire_fire_cost
    )
    return profit - m.d_p_plus + m.d_p_minus == dt.goal_p


model.ProfitGoal = Constraint(rule=profit_goal_rule)


def workforce_change_goal_rule(m):
    total_turnover = sum(m.H[t] + m.L[t] for t in m.T)
    return total_turnover - m.d_cw_plus + m.d_cw_minus == dt.goal_cw


model.WorkforceChangeGoal = Constraint(rule=workforce_change_goal_rule)


def delay_goal_rule(m):
    total_delay = sum(m.pf[i, t] for i in m.I for t in m.T)
    return total_delay - m.d_id_plus + m.d_id_minus == dt.goal_id


model.DelayGoal = Constraint(rule=delay_goal_rule)

# ----------------------------------------------------------------------------
# 6. Hard constraints
# ----------------------------------------------------------------------------


# 3-3-3-1 protein raw-material warehouse capacity
def protein_storage_rule(m, t):
    return sum(inv_prev(m, j, t) + m.X[j, t] for j in PROTEIN_MATERIALS) <= dt.U_Protein


model.ProteinStorage = Constraint(model.T, rule=protein_storage_rule)


# 3-3-3-2 cold-storage raw-material warehouse capacity
def cold_storage_rule(m, t):
    return sum(inv_prev(m, j, t) + m.X[j, t] for j in COLD_MATERIALS) <= dt.U_Cold


model.ColdStorage = Constraint(model.T, rule=cold_storage_rule)


# 3-3-3-3 dry raw-material warehouse capacity
def dry_storage_rule(m, t):
    return sum(inv_prev(m, j, t) + m.X[j, t] for j in DRY_MATERIALS) <= dt.U_Dry


model.DryStorage = Constraint(model.T, rule=dry_storage_rule)


# 3-3-3-4 raw material inventory balance
def material_balance_rule(m, j, t):
    return inv_prev(m, j, t) + m.X[j, t] - m.Y[j, t] == m.Inv[j, t]


model.MaterialBalance = Constraint(model.J, model.T, rule=material_balance_rule)


# 3-3-3-5 raw material consumption (formulation requirement)
def material_consumption_rule(m, j, t):
    return sum(dt.a_ij[i, j] * (m.Pn[i, t] + m.Po[i, t]) for i in m.I) == m.Y[j, t]


model.MaterialConsumption = Constraint(model.J, model.T, rule=material_consumption_rule)


# 3-3-3-6 normal-time and overtime labor-hour capacity
def normal_time_capacity_rule(m, t):
    return sum(m.Pn[i, t] * dt.wn_it[i] for i in m.I) <= m.W[t] * dt.P_1 * dt.b_t[t]


model.NormalTimeCapacity = Constraint(model.T, rule=normal_time_capacity_rule)


def overtime_capacity_rule(m, t):
    return sum(m.Po[i, t] * dt.wn_it[i] for i in m.I) <= m.W[t] * dt.P_2 * dt.b_t[t]


model.OvertimeCapacity = Constraint(model.T, rule=overtime_capacity_rule)


# 3-3-3-7 workforce balance
def workforce_balance_rule(m, t):
    return m.W[t] == w_prev(m, t) + m.H[t] - m.L[t]


model.WorkforceBalance = Constraint(model.T, rule=workforce_balance_rule)


# 3-3-3-8 order fulfilment / finished-goods inventory balance, and sales definition
def demand_balance_rule(m, i, t):
    return (
        sf_prev(m, i, t) + m.Pn[i, t] + m.Po[i, t] + m.pf[i, t]
        == dt.d_it[i, t] + m.sf[i, t] + pf_prev(m, i, t)
    )


model.DemandBalance = Constraint(model.I, model.T, rule=demand_balance_rule)


def sales_definition_rule(m, i, t):
    return m.fi[i, t] == dt.d_it[i, t] - (m.pf[i, t] - pf_prev(m, i, t))


model.SalesDefinition = Constraint(model.I, model.T, rule=sales_definition_rule)


# 3-3-3-9 maximum available workforce
def max_workforce_rule(m, t):
    return m.W[t] <= dt.W_max_t[t]


model.MaxWorkforce = Constraint(model.T, rule=max_workforce_rule)


# 3-3-3-10 protein raw-material supplier capacity
def protein_supply_rule(m, j, t):
    return m.X[j, t] <= dt.Sup_jt[j, t]


model.ProteinSupply = Constraint(PROTEIN_MATERIALS, model.T, rule=protein_supply_rule)


# 3-3-3-11 protein raw-material shelf-life (must be used within 3 weeks)
def protein_shelf_life_rule(m, j, t):
    if t < 4:
        return Constraint.Skip
    return m.Inv[j, t - 3] <= sum(m.Y[j, k] for k in range(t - 2, t + 1))


model.ProteinShelfLife = Constraint(PROTEIN_MATERIALS, model.T, rule=protein_shelf_life_rule)


# 3-3-3-12 finished-product shelf-life (must be sold within 2 weeks)
def product_shelf_life_rule(m, i, t):
    if t < 3:
        return Constraint.Skip
    return m.sf[i, t - 2] <= sum(m.fi[i, k] for k in range(t - 1, t + 1))


model.ProductShelfLife = Constraint(model.I, model.T, rule=product_shelf_life_rule)


# 3-3-3-13 production / setup linking constraint
def setup_link_rule(m, i, t):
    return m.Pn[i, t] + m.Po[i, t] <= dt.M * m.y[i, t]


model.SetupLink = Constraint(model.I, model.T, rule=setup_link_rule)


# 3-3-3-14 cutter room capacity
def cutter_capacity_rule(m, t):
    return sum(
        dt.tc_i[i] * (m.Pn[i, t] + m.Po[i, t]) + dt.SUc_i[i] * m.y[i, t] for i in m.I
    ) <= dt.Cc_t[t]


model.CutterCapacity = Constraint(model.T, rule=cutter_capacity_rule)


# 3-3-3-15 cooking room capacity
def cooking_capacity_rule(m, t):
    return sum(
        dt.to_i[i] * (m.Pn[i, t] + m.Po[i, t]) + dt.SUo_i[i] * m.y[i, t] for i in m.I
    ) <= dt.Co_t[t]


model.CookingCapacity = Constraint(model.T, rule=cooking_capacity_rule)


# 3-3-3-16 packaging capacity
def packaging_capacity_rule(m, t):
    return sum(dt.tp_i[i] * (m.Pn[i, t] + m.Po[i, t]) for i in m.I) <= dt.Cp_t[t]


model.PackagingCapacity = Constraint(model.T, rule=packaging_capacity_rule)

# ----------------------------------------------------------------------------
# 7. Main goal-programming objective function
# ----------------------------------------------------------------------------


def main_objective_rule(m):
    return dt.w1 * m.d_p_minus + dt.w2 * m.d_cw_plus + dt.w3 * m.d_id_plus


model.OBJ = Objective(rule=main_objective_rule, sense=minimize)

# ----------------------------------------------------------------------------
# 8. Solve with GLPK
# ----------------------------------------------------------------------------
solver = SolverFactory("glpk", executable=GLPK_PATH)
results = solver.solve(model, tee=True)

solver_status = str(results.solver.status)
termination_condition = str(results.solver.termination_condition)

print("\n" + "=" * 78)
print(f"Solver status        : {solver_status}")
print(f"Termination condition: {termination_condition}")
print("=" * 78)

if termination_condition not in ("optimal", "feasible", "locallyOptimal"):
    print("The solver did not report an optimal solution. Results printed below")
    print("(if any) may not be reliable. Please check the model / data.")

# ----------------------------------------------------------------------------
# 9. Console output helpers
# ----------------------------------------------------------------------------


def print_two_index_table(title, var, rows, cols, row_label, col_label, width=9):
    print("\n" + "-" * 78)
    print(title)
    print("-" * 78)
    header = f"{row_label:>6}" + "".join(f"{col_label}{c:<{width - len(col_label)}}" for c in cols)
    print(header)
    for r in rows:
        line = f"{r:>6}"
        for c in cols:
            val = value(var[r, c])
            line += f"{val:>{width}.2f}"
        print(line)


def print_one_index_table(title, vars_dict, cols, col_label, width=9):
    print("\n" + "-" * 78)
    print(title)
    print("-" * 78)
    header = f"{'var':>8}" + "".join(f"{col_label}{c:<{width - len(col_label)}}" for c in cols)
    print(header)
    for name, var in vars_dict.items():
        line = f"{name:>8}"
        for c in cols:
            val = value(var[c])
            line += f"{val:>{width}.2f}"
        print(line)


# --- Two-index decision variable tables (products x weeks) ------------------
print_two_index_table("fi_it : sales quantity of product i in week t (tons)",
                       model.fi, PRODUCTS, WEEKS, "i", "t")
print_two_index_table("Pn_it : normal-time production of product i in week t (tons)",
                       model.Pn, PRODUCTS, WEEKS, "i", "t")
print_two_index_table("Po_it : overtime production of product i in week t (tons)",
                       model.Po, PRODUCTS, WEEKS, "i", "t")
print_two_index_table("sf_it : end-of-week finished-product inventory (tons)",
                       model.sf, PRODUCTS, WEEKS, "i", "t")
print_two_index_table("pf_it : delivery delay (backorder) of product i (tons)",
                       model.pf, PRODUCTS, WEEKS, "i", "t")
print_two_index_table("y_it  : 1 if product i is produced in week t",
                       model.y, PRODUCTS, WEEKS, "i", "t", width=5)

# --- Two-index decision variable tables (materials x weeks) -----------------
print_two_index_table("X_jt : raw material j purchased in week t (tons)",
                       model.X, MATERIALS, WEEKS, "j", "t")
print_two_index_table("Y_jt : raw material j consumed in week t (tons)",
                       model.Y, MATERIALS, WEEKS, "j", "t")
print_two_index_table("Inv_jt : end-of-week raw material inventory (tons)",
                       model.Inv, MATERIALS, WEEKS, "j", "t")

# --- One-index workforce variables -------------------------------------------
print_one_index_table("Workforce variables (people)",
                       {"W_t": model.W, "H_t": model.H, "L_t": model.L},
                       WEEKS, "t")

# --- Deviation variables and objective value ---------------------------------
print("\n" + "=" * 78)
print("Deviation (auxiliary) variables")
print("=" * 78)
print(f"d_p_plus   (profit overachievement)          = {value(model.d_p_plus):,.2f}")
print(f"d_p_minus  (profit shortfall)                = {value(model.d_p_minus):,.2f}")
print(f"d_cw_plus  (workforce-turnover overachievement) = {value(model.d_cw_plus):,.2f}")
print(f"d_cw_minus (workforce-turnover shortfall)       = {value(model.d_cw_minus):,.2f}")
print(f"d_id_plus  (delivery-delay overachievement)     = {value(model.d_id_plus):,.2f}")
print(f"d_id_minus (delivery-delay shortfall)           = {value(model.d_id_minus):,.2f}")

print("\n" + "=" * 78)
print(f"Optimal objective value  Z = {value(model.OBJ):,.6f}")
print("=" * 78)

# ----------------------------------------------------------------------------
# 10. output 
# ----------------------------------------------------------------------------


def write_two_index_excel(filename, var, rows, cols, row_label, col_label, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel sheet-name length limit

    ws.cell(row=1, column=1, value=f"{row_label}\\{col_label}")
    for c_idx, c in enumerate(cols, start=2):
        ws.cell(row=1, column=c_idx, value=f"{col_label}{c}")

    for r_idx, r in enumerate(rows, start=2):
        ws.cell(row=r_idx, column=1, value=f"{row_label}{r}")
        for c_idx, c in enumerate(cols, start=2):
            ws.cell(row=r_idx, column=c_idx, value=round(value(var[r, c]), 4))

    wb.save(filename)
    print(f"Saved: {filename}")


write_two_index_excel("fi_output.xlsx", model.fi, PRODUCTS, WEEKS, "i", "t", "fi_it")
write_two_index_excel("Pn_output.xlsx", model.Pn, PRODUCTS, WEEKS, "i", "t", "Pn_it")
write_two_index_excel("Po_output.xlsx", model.Po, PRODUCTS, WEEKS, "i", "t", "Po_it")


def write_workforce_excel(filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Workforce"
    ws.append(["t", "W_t", "H_t", "L_t"])
    for t in WEEKS:
        ws.append([t, round(value(model.W[t]), 2),
                   round(value(model.H[t]), 2),
                   round(value(model.L[t]), 2)])
    wb.save(filename)
    print(f"Saved: {filename}")


write_workforce_excel("workforce_output.xlsx")

print("\nDone.")
